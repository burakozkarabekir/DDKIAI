"""Copilot kullanım raporları (Excel / CSV).

Desteklenen yapılar
- Kişi bazında: Microsoft 365 yönetim merkezi "Copilot kullanımı" dışa aktarımı (User Principal Name, Display Name,
  Last activity date, Teams/Word/Excel/PowerPoint/Outlook/OneNote/Loop/Copilot Chat last activity date …),
  Viva Insights Copilot panosu çıktıları ya da kendi kullanıcı listeniz (kullanıcı, departman, lisans, aksiyon sayısı …).
- Özet tablo: departman / ay bazında sayılar (lisans, aktif kullanıcı, istem sayısı …).

Sütunların rolü başlık adından ve değerlerden otomatik tahmin edilir; Copilot sayfasında değiştirilebilir.
Kişisel veriler (ad, e-posta) yalnızca bu uygulamada kalır: AI'a sadece toplam ve oranlar gönderilir.
"""
import csv
import io
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from . import db
from .knowledge import norm

APPS = ["Teams", "Word", "Excel", "PowerPoint", "Outlook", "OneNote", "Loop", "Copilot Chat"]
_APP_KEYS = [("copilot chat", "Copilot Chat"), ("bing chat", "Copilot Chat"), ("business chat", "Copilot Chat"), ("chat", "Copilot Chat"),
             ("teams", "Teams"), ("word", "Word"), ("excel", "Excel"), ("powerpoint", "PowerPoint"), ("outlook", "Outlook"),
             ("onenote", "OneNote"), ("loop", "Loop")]

ROLES = {
    "user_id": "Kullanıcı kimliği (e-posta / UPN)",
    "user_name": "Ad soyad",
    "department": "Departman / birim",
    "company": "Şirket",
    "license": "Lisans durumu",
    "refresh_date": "Rapor tarihi",
    "last_activity": "Son Copilot etkinliği",
    "period": "Dönem / ay",
    "metric": "Sayısal metrik",
    "ignore": "Kullanma",
}
for _a in APPS:
    ROLES[f"app:{_a}"] = f"{_a} son kullanım / kullanım"

ACTIVE_DAYS = 28  # Microsoft'un "aktif kullanıcı" tanımıyla uyumlu: son 28 gün


# ---------------------------------------------------------------- okuma

def read_tables(path: Path) -> list[dict]:
    """Dönüş: [{"sheet": ad, "header": [...], "rows": [[...], ...]}] — başlık satırı otomatik bulunur."""
    ext = path.suffix.lower()
    raw_sheets = []
    if ext == ".csv":
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t") if text.strip() else csv.excel
        raw_sheets.append(("CSV", list(csv.reader(io.StringIO(text), dialect))))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        for ws in wb.worksheets:
            raw_sheets.append((ws.title, [list(r) for r in ws.iter_rows(values_only=True)]))
    out = []
    for name, rows in raw_sheets:
        rows = [r for r in rows if any(v not in (None, "") for v in r)]
        if len(rows) < 2:
            continue
        hi = _header_index(rows)
        header = [str(v).strip() if v not in (None, "") else f"Sütun {i + 1}" for i, v in enumerate(rows[hi])]
        body = [r + [None] * (len(header) - len(r)) for r in rows[hi + 1:]]
        out.append({"sheet": name, "header": header, "rows": [r[:len(header)] for r in body]})
    return out


def _header_index(rows: list[list]) -> int:
    """İlk 15 satır içinde en çok metin hücresi olan, kendisinden sonra veri gelen satır."""
    best, best_score = 0, -1
    for i, r in enumerate(rows[:15]):
        texts = sum(1 for v in r if isinstance(v, str) and v.strip() and not re.fullmatch(r"[\d.,:/\- ]+", v.strip()))
        if texts > best_score and i < len(rows) - 1:
            best, best_score = i, texts
    return best


def _to_date(v) -> str | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M"):
        try:
            return datetime.strptime(s[:19], fmt).date().isoformat()
        except ValueError:
            continue
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else None


def _to_num(v) -> float | None:
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or "").strip().replace("%", "").replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _cell(v):
    if isinstance(v, (datetime, date)):
        return _to_date(v)
    return v


# ---------------------------------------------------------------- sütun rolleri

def detect_roles(header: list[str], rows: list[list]) -> list[dict]:
    sample = rows[:200]
    cols = []
    used = set()
    for i, h in enumerate(header):
        n = norm(h)
        vals = [r[i] for r in sample if i < len(r) and r[i] not in (None, "")]
        dates = sum(1 for v in vals if _to_date(v))
        nums = sum(1 for v in vals if _to_num(v) is not None)
        is_date = vals and dates >= 0.6 * len(vals)
        is_num = vals and nums >= 0.8 * len(vals) and not is_date
        role = "ignore"
        app = next((a for k, a in _APP_KEYS if k in n), None)
        if any(k in n for k in ("user principal", "upn", "e-posta", "eposta", "email", "e-mail", "mail adresi", "kullanici kimligi",
                                "kullanici adi")) or (n in ("mail", "kullanici") and not is_num):
            role = "user_id"
        elif any(k in n for k in ("display name", "ad soyad", "adi soyadi", "isim", "calisan adi", "personel")) or n in ("ad", "name"):
            role = "user_name"
        elif any(k in n for k in ("refresh", "rapor tarihi", "report date")):
            role = "refresh_date"
        elif app and (is_date or "last" in n or "son" in n or "tarih" in n):
            role = f"app:{app}"
        elif app and is_num:
            role = f"app:{app}"
        elif any(k in n for k in ("departman", "department", "birim", "bolum", "direktorluk", "mudurluk", "organization", "organizasyon",
                                  "fonksiyon", "ekip")):
            role = "department"
        elif any(k in n for k in ("sirket", "company", "firma")):
            role = "company"
        elif any(k in n for k in ("lisans", "license", "lisanslı", "assigned")) and not is_num:
            role = "license"
        elif ("last activity" in n or "son etkinlik" in n or "son aktivite" in n or "son kullanim" in n) and (is_date or not vals):
            role = "last_activity"
        elif any(k in n for k in ("donem", "period", "month")) or n in ("ay", "tarih", "date") or (is_date and "tarih" in n):
            role = "period"
        elif is_num:
            role = "metric"
        if role in ("user_id", "user_name", "department", "company", "refresh_date", "last_activity", "period", "license") and role in used:
            role = "ignore" if role != "department" else "company"
        used.add(role)
        cols.append({"name": h, "role": role})
    return cols


def kind_of(columns: list[dict]) -> str:
    roles = {c["role"] for c in columns}
    return "user" if roles & {"user_id", "user_name"} else "aggregate"


# ---------------------------------------------------------------- içe aktarma

def import_document(doc_id: int) -> list[int]:
    from .storage import path_of
    doc = db.row("SELECT * FROM documents WHERE id = ?", (doc_id,))
    path = path_of(doc["storage_path"]) if doc else None
    if not path or not path.exists():
        return []
    db.execute("DELETE FROM copilot_datasets WHERE document_id = ?", (doc_id,))
    ids = []
    for t in read_tables(path):
        cols = detect_roles(t["header"], t["rows"])
        if not any(c["role"] != "ignore" for c in cols):
            continue
        rows = [{h: _cell(v) for h, v in zip(t["header"], r)} for r in t["rows"]]
        refresh = _refresh_date(cols, rows) or doc["doc_date"] or _date_in(doc["filename"]) or doc["uploaded_at"][:10]
        did = db.insert("copilot_datasets", {
            "document_id": doc_id, "title": doc["title"] + (f" · {t['sheet']}" if t["sheet"] not in ("CSV", "Sheet", "Sheet1", "Sayfa1") else ""),
            "period": refresh[:7], "sheet": t["sheet"], "kind": kind_of(cols), "columns": json.dumps(cols, ensure_ascii=False),
            "row_count": len(rows), "refresh_date": refresh, "created_at": db.now()})
        with db.tx() as c:
            c.executemany("INSERT INTO copilot_rows (dataset_id, data) VALUES (?, ?)",
                          [(did, json.dumps(r, ensure_ascii=False, default=str)) for r in rows])
        ids.append(did)
    return ids


def _refresh_date(cols, rows) -> str | None:
    col = next((c["name"] for c in cols if c["role"] == "refresh_date"), None)
    if col:
        ds = [d for d in (_to_date(r.get(col)) for r in rows) if d]
        if ds:
            return max(ds)
    return None


def _date_in(text: str) -> str | None:
    m = re.search(r"(\d{4})[-_.](\d{2})[-_.](\d{2})|(\d{2})[-_.](\d{2})[-_.](\d{4})", text or "")
    if not m:
        return None
    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m.group(1) else f"{m.group(6)}-{m.group(5)}-{m.group(4)}"


def set_columns(dataset_id: int, roles: dict[str, str]) -> None:
    ds = db.row("SELECT * FROM copilot_datasets WHERE id = ?", (dataset_id,))
    cols = json.loads(ds["columns"])
    for c in cols:
        if roles.get(c["name"]) in ROLES:
            c["role"] = roles[c["name"]]
    rows = rows_of(dataset_id)
    refresh = _refresh_date(cols, rows) or ds["refresh_date"]
    db.update("copilot_datasets", dataset_id, {"columns": json.dumps(cols, ensure_ascii=False), "kind": kind_of(cols),
                                               "refresh_date": refresh, "period": (refresh or "")[:7] or ds["period"]})


def rows_of(dataset_id: int) -> list[dict]:
    return [json.loads(r["data"]) for r in db.rows("SELECT data FROM copilot_rows WHERE dataset_id = ? ORDER BY id", (dataset_id,))]


def datasets() -> list[dict]:
    return db.rows("SELECT * FROM copilot_datasets ORDER BY COALESCE(refresh_date, created_at) DESC, id DESC")


# ---------------------------------------------------------------- analitik

def analyze(dataset_id: int) -> dict:
    ds = db.row("SELECT * FROM copilot_datasets WHERE id = ?", (dataset_id,))
    cols = json.loads(ds["columns"])
    rows = rows_of(dataset_id)
    by_role: dict[str, list[str]] = {}
    for c in cols:
        by_role.setdefault(c["role"], []).append(c["name"])
    return (_analyze_users if ds["kind"] == "user" else _analyze_aggregate)(ds, cols, rows, by_role)


def _activity_dates(r: dict, by_role: dict) -> tuple[str | None, dict[str, str | float | None]]:
    apps = {}
    for role, names in by_role.items():
        if role.startswith("app:"):
            v = r.get(names[0])
            apps[role[4:]] = _to_date(v) or _to_num(v)
    last = _to_date(r.get(by_role["last_activity"][0])) if by_role.get("last_activity") else None
    if not last:
        ds = [v for v in apps.values() if isinstance(v, str)]
        last = max(ds) if ds else None
    return last, apps


def _analyze_users(ds, cols, rows, by_role) -> dict:
    ref = datetime.fromisoformat(ds["refresh_date"]).date() if ds["refresh_date"] else date.today()
    cutoff = (ref - timedelta(days=ACTIVE_DAYS)).isoformat()
    dept_col = (by_role.get("department") or by_role.get("company") or [None])[0]
    lic_col = (by_role.get("license") or [None])[0]
    metric_cols = by_role.get("metric", [])
    total, active, never = 0, 0, 0
    buckets = {"Son 7 gün": 0, "8-28 gün": 0, "29-90 gün": 0, "90+ gün": 0, "Hiç kullanmadı": 0}
    depts: dict[str, dict] = {}
    apps = {a: 0 for a in APPS if f"app:{a}" in by_role}
    metrics = {m: 0.0 for m in metric_cols}
    inactive = []
    for r in rows:
        if lic_col and str(r.get(lic_col) or "").strip().lower() in ("hayır", "hayir", "no", "false", "0", "yok", "unlicensed"):
            continue
        total += 1
        last, app_vals = _activity_dates(r, by_role)
        is_active = bool(last and last >= cutoff)
        if not last and app_vals:  # tarih yok ama sayısal kullanım varsa
            is_active = any(isinstance(v, float) and v > 0 for v in app_vals.values())
        active += is_active
        if last:
            age = (ref - datetime.fromisoformat(last).date()).days
            key = "Son 7 gün" if age <= 7 else "8-28 gün" if age <= 28 else "29-90 gün" if age <= 90 else "90+ gün"
        else:
            key = "Hiç kullanmadı" if not is_active else "Son 7 gün"
            never += not is_active
        buckets[key] += 1
        for a, v in app_vals.items():
            if (isinstance(v, str) and v >= cutoff) or (isinstance(v, float) and v > 0):
                apps[a] = apps.get(a, 0) + 1
        for m in metric_cols:
            metrics[m] += _to_num(r.get(m)) or 0
        d = str(r.get(dept_col) or "Belirtilmemiş").strip() if dept_col else "Belirtilmemiş"
        dd = depts.setdefault(d, {"name": d, "total": 0, "active": 0, **{f"m:{m}": 0.0 for m in metric_cols}})
        dd["total"] += 1
        dd["active"] += is_active
        for m in metric_cols:
            dd[f"m:{m}"] += _to_num(r.get(m)) or 0
        if not is_active:
            name = r.get((by_role.get("user_name") or by_role.get("user_id") or [""])[0])
            inactive.append({"name": name, "dept": d, "last": last})
    for d in depts.values():
        d["pct"] = round(100 * d["active"] / d["total"]) if d["total"] else 0
    return {
        "kind": "user", "refresh_date": ds["refresh_date"], "total": total, "active": active,
        "adoption": round(100 * active / total, 1) if total else 0, "inactive": total - active, "never": never,
        "buckets": buckets, "apps": sorted(({"app": a, "users": n, "pct": round(100 * n / total) if total else 0} for a, n in apps.items()),
                                          key=lambda x: -x["users"]),
        "departments": sorted(depts.values(), key=lambda d: (-d["total"], d["name"])), "metrics": metrics,
        "inactive_list": sorted(inactive, key=lambda x: (x["last"] or ""))[:300], "has_dept": bool(dept_col),
    }


def _analyze_aggregate(ds, cols, rows, by_role) -> dict:
    dim = (by_role.get("department") or by_role.get("company") or by_role.get("period") or [None])[0]
    metric_cols = by_role.get("metric", []) + [n for r, ns in by_role.items() if r.startswith("app:") for n in ns]
    groups: dict[str, dict] = {}
    totals = {m: 0.0 for m in metric_cols}
    for r in rows:
        key = str(r.get(dim) or "Toplam").strip() if dim else "Toplam"
        if norm(key) in ("toplam", "total", "genel toplam", "grand total") and dim:
            continue  # tablodaki toplam satırını tekrar sayma
        g = groups.setdefault(key, {"name": key, **{m: 0.0 for m in metric_cols}})
        for m in metric_cols:
            v = _to_num(r.get(m)) or 0
            g[m] += v
            totals[m] += v
    return {"kind": "aggregate", "refresh_date": ds["refresh_date"], "dimension": dim, "metrics": metric_cols,
            "groups": sorted(groups.values(), key=lambda g: -(g[metric_cols[0]] if metric_cols else 0)), "totals": totals,
            "rows": len(rows)}


def trend() -> list[dict]:
    """Kişi bazındaki veri setlerinden dönem dönem aktif kullanıcı ve benimseme oranı."""
    out = {}
    for ds in reversed(datasets()):
        if ds["kind"] != "user":
            continue
        a = analyze(ds["id"])
        out[ds["period"] or (ds["refresh_date"] or ds["created_at"])[:7]] = {"period": ds["period"], "active": a["active"], "total": a["total"],
                                                       "adoption": a["adoption"], "dataset_id": ds["id"]}
    return list(out.values())


def headline() -> dict | None:
    """Ana sayfa ve AI Security özet kartı için son kişi bazlı veri setinin özeti."""
    t = trend()
    if not t:
        return None
    last = t[-1]
    prev = t[-2] if len(t) > 1 else None
    return {**last, "delta": round(last["adoption"] - prev["adoption"], 1) if prev else None, "points": t}


def ai_summary() -> str:
    """AI'a gönderilecek kişisel veri içermeyen özet."""
    ds = datasets()
    if not ds:
        return "# Copilot kullanımı\nVeri yüklenmedi."
    lines = ["# Copilot kullanımı (kişisel veri içermez; yalnızca toplamlar)"]
    for t in trend():
        lines.append(f"- {t['period']}: {t['active']}/{t['total']} aktif (%{t['adoption']})")
    latest = next((d for d in ds if d["kind"] == "user"), None)
    if latest:
        a = analyze(latest["id"])
        lines.append(f"\nSon rapor ({a['refresh_date']}): {a['total']} lisanslı, {a['active']} aktif (son {ACTIVE_DAYS} gün), "
                     f"%{a['adoption']} benimseme, {a['never']} kişi hiç kullanmamış.")
        lines.append("Uygulama bazında aktif kullanıcı: " + ", ".join(f"{x['app']} {x['users']} (%{x['pct']})" for x in a["apps"]))
        lines.append("Son etkinlik dağılımı: " + ", ".join(f"{k}: {v}" for k, v in a["buckets"].items()))
        if a["has_dept"]:
            lines.append("Departman bazında (lisans / aktif / %):")
            lines += [f"  - {d['name']}: {d['total']} / {d['active']} / %{d['pct']}" for d in a["departments"][:40]]
        if a["metrics"]:
            lines.append("Toplam metrikler: " + ", ".join(f"{k}: {v:,.0f}" for k, v in a["metrics"].items()))
    for d in ds:
        if d["kind"] == "aggregate":
            a = analyze(d["id"])
            lines.append(f"\nÖzet tablo '{d['title']}' ({d['refresh_date']}), boyut: {a['dimension'] or '—'}")
            lines.append("  Toplamlar: " + ", ".join(f"{k}: {v:,.0f}" for k, v in a["totals"].items()))
            lines += [f"  - {g['name']}: " + ", ".join(f"{m} {g[m]:,.0f}" for m in a["metrics"][:6]) for g in a["groups"][:25]]
            break
    return "\n".join(lines)


PROMPT = """Aşağıda şirketin Microsoft 365 Copilot kullanım verisinin özeti var (kişisel veri içermez).
Markdown, şu başlıklarla analiz et:
## Genel tablo (benimseme oranı, aktif kullanıcı, dönemsel eğilim — 3-4 cümle)
## Departman bazında benimseme (öne çıkanlar, geride kalanlar; lisans başına aktif kullanım)
## Uygulama bazında kullanım (hangi uygulamalarda değer üretiliyor, hangilerinde kullanılmıyor)
## Lisans verimliliği (kullanılmayan lisanslar, yeniden dağıtım önerisi, tahmini tasarruf yaklaşımı)
## Önerilen aksiyonlar (eğitim, şampiyon ağı, kullanım senaryoları; departman bazında, en fazla 6)
Rakam uydurma; yalnızca verilen sayıları kullan."""


def ai_analysis() -> str:
    from . import ai
    summary = ai_summary()
    if ai.is_live():
        return ai.complete([{"role": "user", "content": f"{PROMPT}\n\n{summary}"}], max_tokens=3000)
    return "> *AI bağlı değil: yalnızca hesaplanan özet gösteriliyor.*\n\n" + summary.replace("# Copilot kullanımı", "## Copilot kullanımı")
