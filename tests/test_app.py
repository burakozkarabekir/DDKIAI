"""Uçtan uca testler: python -m pytest -q

Uygulama boş veritabanıyla başlar (demo veri yoktur); testler kendi verisini arayüz uçlarından oluşturur.
"""
import io
import json
import os
import tempfile
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer

os.environ["DATA_DIR"] = tempfile.mkdtemp()
os.environ["AI_PROVIDER"] = "mock"
os.environ["AUTH_ENABLED"] = "0"

import docx  # noqa: E402
import openpyxl  # noqa: E402
import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import ai, config, copilot, db, domain, knowledge, org, pipeline, seed, status  # noqa: E402
from app.main import _f, _md, app  # noqa: E402

PAGES = ["/", "/projeler", "/projeler?view=liste", "/projeler/yeni", "/durum-girisi", "/zaman-cizelgesi", "/zaman-cizelgesi?grup=asama",
         "/dokumanlar", "/raporlar", "/inceleme", "/aksiyonlar", "/danisman", "/guvenlik", "/guvenlik/copilot",
         "/alan/yonetici-raporlari", "/alan/ai-yonetisimi", "/surec", "/ayarlar", "/yonetim"]


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def _docx(lines) -> bytes:
    d = docx.Document()
    d.add_heading(lines[0], level=0)
    for line in lines[1:]:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _brand(name) -> int:
    return db.scalar("SELECT id FROM brands WHERE name = ?", (name,))


def _dept(name) -> int:
    return db.scalar("SELECT id FROM departments WHERE name = ?", (name,))


def _create(client, code, name, stage="gelistirme", brands=(), **extra) -> int:
    data = {"code": code, "name": name, "stage": stage, "health": "yesil", "priority": "Orta", "measurement_status": "yok",
            "security_review": "yok", **extra, "brand_ids": [str(_brand(b)) for b in brands]}
    r = client.post("/projeler", data=data, follow_redirects=False)
    assert r.status_code == 303, r.text
    return int(r.headers["location"].rsplit("/", 1)[1])


# ---------------------------------------------------------------- boş kurulum, organizasyon

def test_fresh_install_has_no_dummy_data(client):
    assert db.scalar("SELECT COUNT(*) FROM projects") == 0
    assert db.scalar("SELECT COUNT(*) FROM documents") == 0
    assert db.scalar("SELECT COUNT(*) FROM initiatives") == 0
    assert db.scalar("SELECT COUNT(*) FROM security_metrics") == 0
    for url in PAGES:
        assert client.get(url).status_code == 200, url
    assert "hero-steps" in client.get("/").text  # boş ana sayfada başlangıç adımları
    assert client.get("/giris", follow_redirects=False).status_code == 303  # giriş kapalı


def test_org_has_ten_departments_and_separate_brands():
    names = [d["name"] for d in org.departments()]
    assert names == [n for n, _ in domain.ORG] and len(names) == 10
    assert names[1] == "VW Ticari Araç - Scania - Thermo King - Meiller - Wielton"
    scania = db.row("SELECT * FROM brands WHERE name = 'Scania'")
    assert scania["department_id"] == _dept(names[1])
    assert db.row("SELECT * FROM brands WHERE name = 'Cupra'")["department_id"] == _dept("Seat - Cupra - Porsche - DOD")
    assert sorted(domain.match_brands("Scania kamyon ve Thermo King")) == ["Scania", "Thermo King"]
    assert "VW Binek" in domain.match_brands("Volkswagen Binek satış")


def test_project_create_edit_with_brands(client):
    pid = _create(client, "AI-007", "Fatura & İrsaliye Eşleştirme", brands=["Scania", "Meiller"])
    p = db.row("SELECT * FROM projects WHERE id = ?", (pid,))
    assert p["code"] == "AI-007" and p["department_id"] == _dept(domain.ORG[1][0])  # departman markadan geldi
    assert sorted(org.brands_by_project()[pid]) == ["Meiller", "Scania"]
    # düzenleme: marka değişir, departman elle verilirse korunur
    client.post(f"/projeler/{pid}", data={"name": "Fatura & İrsaliye Eşleştirme", "code": "AI-007", "brand_ids": [str(_brand("Audi"))],
                                          "department_id": str(_dept("Mali İşler"))})
    p = db.row("SELECT * FROM projects WHERE id = ?", (pid,))
    assert p["department_id"] == _dept("Mali İşler") and org.brands_by_project()[pid] == ["Audi"]
    assert "Audi" in client.get(f"/projeler?view=liste&marka={_brand('Audi')}").text
    assert client.get(f"/projeler/{pid}").status_code == 200
    _create(client, "AI-003", "Yedek Parça Talep Tahmini", brands=["Yedek Parça ve Lojistik"] if _brand("Yedek Parça ve Lojistik") else [],
            department_id=str(_dept("Yedek Parça ve Lojistik")), target_date="2026-11-30")
    _create(client, "AI-006", "İK Asistanı", stage="fikir", department_id=str(_dept("İnsan Kaynakları ve Süreç Yönetimi")))
    html = client.get("/").text
    assert "Departman × aşama" in html and "Marka kırılımı" in html


def test_manual_status_update_and_bulk(client):
    pid = db.scalar("SELECT id FROM projects WHERE code = 'AI-003'")
    r = client.post(f"/projeler/{pid}/durum", data={"health": "sari", "stage": "pilot", "progress": "60", "summary": "Pilot başladı",
                                                     "blockers": "ERP erişimi", "next_steps": "Kullanıcı eğitimi"}, follow_redirects=False)
    assert r.status_code == 303
    p = db.row("SELECT * FROM projects WHERE id = ?", (pid,))
    assert (p["health"], p["stage"]) == ("sari", "pilot") and p["last_status_at"]
    last = status.latest(pid)
    assert last["progress"] == 60 and last["blockers"] == "ERP erişimi"
    assert db.scalar("SELECT 1 FROM stage_history WHERE project_id = ? AND to_stage = 'pilot'", (pid,))
    assert "Pilot başladı" in client.get(f"/projeler/{pid}?tab=bilgi").text
    # toplu giriş: yalnız değişen satır kaydedilir
    other = db.scalar("SELECT id FROM projects WHERE code = 'AI-007'")
    n = db.scalar("SELECT COUNT(*) FROM status_updates")
    client.post("/durum-girisi", data={f"p_{other}_health": "kirmizi", f"p_{other}_summary": "Tedarikçi gecikmesi",
                                       f"p_{pid}_health": "sari", f"p_{pid}_stage": "pilot"})
    assert db.scalar("SELECT COUNT(*) FROM status_updates") == n + 1
    assert db.scalar("SELECT health FROM projects WHERE id = ?", (other,)) == "kirmizi"
    assert "Tedarikçi gecikmesi" in ai.portfolio_context()  # manuel güncellemeler AI bağlamına girer


# ---------------------------------------------------------------- dokümanlar

def test_upload_word_without_project_links_automatically(client):
    data = _docx(["Komite notu", "AI-007 Fatura & İrsaliye Eşleştirme: Karar: UAT 2 hafta sonra başlayacak.",
                  "Aksiyon: Muhasebe test kullanıcılarını belirlesin."])
    r = client.post("/dokumanlar", files=[("files", ("Komite Notu.docx", data, "application/octet-stream"))],
                    data={"doc_type": "otomatik"}, follow_redirects=False)
    assert r.status_code == 303
    did = int(r.headers["location"].rsplit("/", 1)[1])
    pipeline.drain()
    d = db.row("SELECT * FROM documents WHERE id = ?", (did,))
    assert d["status"] == "islendi" and d["doc_type"] == "toplanti"
    assert db.scalar("SELECT 1 FROM document_links l JOIN projects p ON p.id = l.project_id WHERE l.document_id = ? AND p.code = 'AI-007'", (did,))
    assert db.scalar("SELECT COUNT(*) FROM facts WHERE document_id = ? AND fact_type IN ('karar', 'aksiyon')", (did,)) == 2
    assert client.get(f"/dokumanlar/{did}").status_code == 200
    # aynı dosya tekrar yüklenince kaydedilmez
    r = client.post("/dokumanlar", files=[("files", ("kopya.docx", data, "application/octet-stream"))], follow_redirects=True)
    assert "zaten" in r.text


def test_multi_upload_and_bad_extension(client):
    files = [("files", ("a.txt", b"AI-003 Yedek Parca: hedef canli tarihi 15.12.2026, saglik: sari.", "text/plain")),
             ("files", ("b.exe", b"x", "application/octet-stream"))]
    r = client.post("/dokumanlar", files=files, data={"doc_type": "dt_durum"}, follow_redirects=True)
    assert "1 doküman yüklendi" in r.text and "desteklenmiyor" in r.text
    pipeline.drain()


def test_review_apply_updates_project(client):
    pending = knowledge.pending()
    assert pending
    f = pending[0]
    client.post("/inceleme", data={"fact_id": str(f["id"]), "action": "onayla"})
    assert db.scalar("SELECT status FROM facts WHERE id = ?", (f["id"],)) == "onaylandi"


def test_exec_report_area_and_analysis(client):
    data = _docx(["Yönetici Özeti Ekim 2026", "AI portföyünde 3 proje ilerliyor. AI-003 pilotta, AI-007 için tedarikçi riski var."])
    r = client.post("/dokumanlar", files=[("files", ("Yonetici Ozeti Ekim.docx", data, "application/octet-stream"))],
                    data={"doc_type": "yonetici_ozeti"}, follow_redirects=False)
    assert r.status_code == 303
    pipeline.drain()
    html = client.get("/alan/yonetici-raporlari").text
    assert "Yonetici Ozeti Ekim" in html or "Yönetici Özeti Ekim" in html
    client.post("/alan/yonetici-raporlari/analiz")
    pipeline.drain()
    assert db.scalar("SELECT 1 FROM ai_insights WHERE kind = 'area:yonetici-raporlari'")
    assert client.get("/alan/yonetici-raporlari").status_code == 200
    assert client.get("/alan/olmayan").status_code == 404


def test_stage_gate_warns(client):
    pid = db.scalar("SELECT id FROM projects WHERE code = 'AI-006'")
    r = client.post(f"/projeler/{pid}/asama", data={"to_stage": "onay"}, follow_redirects=True)
    assert "kapı kriterleri eksik" in r.text


def test_departments_and_brands_crud(client):
    client.post("/ayarlar/departman", data={"name": "Test Departmanı", "description": "geçici"})
    dep = db.row("SELECT * FROM departments WHERE name = 'Test Departmanı'")
    assert dep
    client.post("/ayarlar/marka", data={"name": "Test Marka", "department_id": str(dep["id"]), "active": "1"})
    bid = _brand("Test Marka")
    assert db.scalar("SELECT department_id FROM brands WHERE id = ?", (bid,)) == dep["id"]
    client.post(f"/ayarlar/marka/{bid}/sil")
    assert not _brand("Test Marka")
    client.post(f"/ayarlar/departman/{dep['id']}/sil")
    assert not db.row("SELECT * FROM departments WHERE id = ?", (dep["id"],))


def test_synthesis_and_advisor_with_page_context(client):
    pid = db.scalar("SELECT id FROM projects WHERE code = 'AI-003'")
    client.post(f"/projeler/{pid}/sentez")
    pipeline.drain()
    assert db.scalar("SELECT 1 FROM ai_insights WHERE kind = 'project' AND project_id = ?", (pid,))
    r = client.post("/api/danisman", json={"question": "Riskli projeler hangileri?", "page": f"/projeler/{pid}", "project_id": pid})
    assert r.status_code == 200 and r.json()["answer_html"]
    hist = client.get("/api/danisman/gecmis").json()["messages"]
    assert hist[-2]["text"] == "Riskli projeler hangileri?" and hist[-1]["html"]
    assert "sohbet balonu" not in client.get("/danisman").text and "fab-chat" not in client.get("/danisman").text
    assert "fab-chat" in client.get("/").text
    assert client.post("/api/danisman/temizle").json()["ok"]
    assert client.get("/api/danisman/gecmis").json()["messages"] == []


# ---------------------------------------------------------------- Copilot

def _copilot_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Report Refresh Date", "User Principal Name", "Display Name", "Department",
               "Last Activity Date", "Microsoft Teams Copilot Last Activity Date", "Word Copilot Last Activity Date"])
    for i in range(10):
        active = i < 4
        ws.append(["2026-10-20", f"kisi{i}@dogus.com", f"Kişi {i}", "Mali İşler" if i % 2 else "Hukuk",
                   "2026-10-15" if active else ("2026-05-01" if i < 8 else None), "2026-10-15" if active else None,
                   "2026-10-10" if i < 2 else None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_copilot_upload_and_analysis(client):
    r = client.post("/guvenlik/copilot/yukle", files={"file": ("Copilot_kullanim_2026-10.xlsx", _copilot_xlsx())}, follow_redirects=False)
    assert r.status_code == 303 and "ds=" in r.headers["location"]
    ds = copilot.datasets()[0]
    assert ds["kind"] == "user" and ds["row_count"] == 10
    a = copilot.analyze(ds["id"])
    assert (a["total"], a["active"], a["adoption"]) == (10, 4, 40.0)
    assert {d["name"] for d in a["departments"]} == {"Mali İşler", "Hukuk"}
    assert any(x["app"] == "Teams" and x["users"] == 4 for x in a["apps"])
    summary = copilot.ai_summary()
    assert "kisi1@" not in summary and "Kişi 1" not in summary  # kişisel veri AI'a gitmez
    assert copilot.headline()["adoption"] == 40.0
    assert "%40" in client.get("/guvenlik/copilot").text
    assert "Copilot benimseme" in client.get("/").text and "Copilot" in client.get("/guvenlik").text
    client.post("/guvenlik/copilot/analiz")
    pipeline.drain()
    assert client.get("/guvenlik/copilot").status_code == 200
    # Copilot dosyası doküman olarak metin indekslemesine girmez
    assert not db.scalar("SELECT text_content FROM documents WHERE id = ?", (ds["document_id"],))
    r = client.post(f"/guvenlik/copilot/{ds['id']}/sil", follow_redirects=False)
    assert r.status_code == 303 and not copilot.datasets()


# ---------------------------------------------------------------- eski sürümden geçiş

def test_legacy_demo_purge_and_department_migration():
    from app.legacy_demo import DEMO_PROJECTS
    code, name = DEMO_PROJECTS[0]
    ts = db.now()
    keep = db.insert("projects", {"code": "UC999", "name": "Kullanıcı projesi", "stage": "fikir", "created_at": ts, "updated_at": ts,
                                  "brands": "Porsche"})
    db.insert("projects", {"code": code, "name": name, "stage": "fikir", "created_at": ts, "updated_at": ts})
    old = db.insert("departments", {"name": "Finans", "sort": 99, "active": 1})
    db.execute("UPDATE projects SET department_id = ? WHERE id = ?", (old, keep))
    demo_user = db.insert("users", {"email": "dt@demo.local", "name": "DT", "role": "dt", "password_hash": "x", "created_at": ts})
    db.log(demo_user, "eski_kayit")  # başka tabloda adı geçen demo kullanıcısı silinemez, pasif yapılır
    db.set_setting("demo_purged_v3", "")
    db.set_setting("org_v3", "")
    seed.purge_legacy_demo()
    assert db.scalar("SELECT active FROM users WHERE id = ?", (demo_user,)) == 0
    org.ensure_org()
    assert not db.scalar("SELECT 1 FROM projects WHERE code = ? AND name = ?", (code, name))
    p = db.row("SELECT * FROM projects WHERE id = ?", (keep,))
    assert p["department_id"] == _dept("Mali İşler")
    assert org.brands_by_project()[keep] == ["Porsche"]
    assert not _dept("Finans") and len(org.departments()) == 10


def test_origin_check(client):
    r = client.post("/yorum", data={"target_type": "project", "target_id": 1, "body": "x"}, headers={"origin": "http://evil.example"})
    assert r.status_code == 403


def test_auth_can_be_reenabled(client):
    config.AUTH_ENABLED = True
    try:
        assert client.get("/", follow_redirects=False).headers["location"].startswith("/giris")
        r = client.post("/giris", data={"email": "admin@demo.local", "password": "Admin123!"}, follow_redirects=False)
        assert r.status_code == 303
        assert client.get("/").status_code == 200
    finally:
        config.AUTH_ENABLED = False
        client.get("/cikis")


def test_live_ai_provider_with_fake_server(client):
    """OpenAI uyumlu uç nokta: kimlik başlığı, JSON modu ve çıkarım sonucunun işlenmesi."""
    seen = []

    class H(BaseHTTPRequestHandler):
        def do_POST(self):
            body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            seen.append((self.path, self.headers.get("Authorization")))
            if body.get("response_format"):
                content = json.dumps({"doc_type": "kazanim", "doc_date": "2026-09-30", "summary": "AI-006 kazanım özeti",
                                      "projects": [{"code": "AI-006", "confidence": 0.95}], "new_projects": [{"name": "Bayi Chatbot"}],
                                      "facts": [{"project_code": "AI-006", "type": "kazanim", "field": "realized_value", "value": None,
                                                 "value_num": 1100000, "unit": "TL", "as_of": "2026-09-30", "confidence": 0.9,
                                                 "evidence": "Gerçekleşen 1,1 milyon TL"}]})
            else:
                content = "Yanıt [D1]"
            out = json.dumps({"choices": [{"message": {"content": content}}]}).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(out)

        def log_message(self, *a):
            pass

    srv = HTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    old = (config.AI_PROVIDER, config.AI_BASE_URL, config.AI_API_KEY, config.AI_MODEL)
    config.AI_PROVIDER, config.AI_BASE_URL, config.AI_API_KEY, config.AI_MODEL = "openai", f"http://127.0.0.1:{srv.server_port}/v1", "k", "m"
    try:
        assert ai.test_connection()["ok"]
        r = client.post("/dokumanlar", files=[("files", ("ik.txt", "İK asistanı kazanım raporu".encode(), "text/plain"))], follow_redirects=False)
        did = int(r.headers["location"].rsplit("/", 1)[1])
        pipeline.drain()
        d = db.row("SELECT * FROM documents WHERE id = ?", (did,))
        assert d["doc_type"] == "kazanim" and d["type_source"] == "ai" and d["doc_date"] == "2026-09-30"
        f = db.row("SELECT * FROM facts WHERE document_id = ? AND fact_type = 'kazanim'", (did,))
        assert f["value_num"] == 1_100_000 and f["status"] == "oneri"
        assert db.scalar("SELECT 1 FROM facts WHERE document_id = ? AND field = 'yeni_proje'", (did,))
        assert seen[0] == ("/v1/chat/completions", "Bearer k")
    finally:
        config.AI_PROVIDER, config.AI_BASE_URL, config.AI_API_KEY, config.AI_MODEL = old
        srv.shutdown()


def test_helpers():
    assert _f("1.500.000") == 1_500_000 and _f("12,5") == 12.5 and _f("") is None
    html = str(_md("<script>x</script> [a](javascript:alert(1)) kaynak [D3]"))
    assert "<script>" not in html and "javascript:" not in html and 'href="/dokumanlar/3"' in html
    assert ai.parse_number("14,5 milyon TL") == 14_500_000 and ai.parse_number("15.200.000 TL") == 15_200_000
    assert knowledge.norm("KPI İLK") == "kpi ilk"


# ---------------------------------------------------------------- v2.1

TRELLO = {
    "name": "AI Impact",
    "lists": [{"id": "L1", "name": "İŞ LİSTESİ"}, {"id": "L2", "name": "INPROGRESS"}, {"id": "L3", "name": "DONE/MAINTENANCE"},
              {"id": "L4", "name": "AI Champion"}, {"id": "L5", "name": "ÖN DEĞERLENDİRME"}],
    "labels": [{"id": "B1", "name": "AI Impact"}, {"id": "B2", "name": "Dependent"}],
    "members": [{"id": "M1", "fullName": "Burak Özkarabekir"}],
    "checklists": [{"id": "K1", "name": "Checklist", "checkItems": [{"name": "DT ön görüşme", "state": "incomplete"}]}],
    "cards": [
        {"id": "68000000a1", "name": "UC88-YP Lojistik-Spare Parts Demand Forecasting", "idList": "L2", "idLabels": ["B1", "B2"],
         "idMembers": ["M1"], "due": "2027-01-08T09:00:00.000Z", "desc": "Yedek parça talep tahmini. Detay.", "idChecklists": ["K1"]},
        {"id": "68000000a2", "name": "UC88- Yedek Parça Lojistik - Stok Optimizasyon", "idList": "L1", "idLabels": [], "idMembers": []},
        {"id": "68000000a3", "name": "UC57-VW Binek - Price Positioning Faz2", "idList": "L3", "idLabels": ["B1"], "idMembers": []},
        {"id": "68000000a4", "name": "UC62 - Micro-Segment & Personalization Engine", "idList": "L5", "idLabels": [], "idMembers": []},
        {"id": "68000000a5", "name": "AI Championlar", "idList": "L4", "idLabels": [], "idMembers": []},
    ],
    "actions": [
        {"type": "updateCard", "date": "2026-10-05T10:39:56.154Z", "data": {"card": {"id": "68000000a1"},
         "listBefore": {"id": "L1", "name": "İŞ LİSTESİ"}, "listAfter": {"id": "L2", "name": "INPROGRESS"}}},
        {"type": "updateCard", "date": "2026-04-06T10:00:00.000Z", "data": {"card": {"id": "68000000a3"},
         "listBefore": {"id": "L2", "name": "INPROGRESS"}, "listAfter": {"id": "L3", "name": "DONE/MAINTENANCE"}}},
        {"type": "commentCard", "date": "2026-09-01T10:00:00.000Z", "data": {"card": {"id": "68000000a1"}, "text": "ERP onayı bekleniyor"},
         "memberCreator": {"fullName": "Burak Özkarabekir"}},
    ],
}


def test_trello_parse_and_import(client):
    from app import trello
    assert trello.parse_name("UC88-YP Lojistik-Spare Parts Demand Forecasting")["unit"] == "YP Lojistik"
    assert trello.parse_name("UC62 - Micro-Segment & Personalization Engine")["unit"] is None
    assert trello.parse_name("UC32 & UC43- Total Cost of Ownership")["extra"] == ["UC43"]
    assert trello.guess_stage("ÖN DEĞERLENDİRME") == "analiz" and trello.guess_stage("AI Champion") == "atla"
    raw = json.dumps(TRELLO)
    r = client.post("/ayarlar/trello/onizle", files={"file": ("t.json", raw.encode(), "application/json")})
    assert r.status_code == 200 and "UC88" in r.text  # tekrar eden kod uyarısı
    yp = _dept("Yedek Parça ve Lojistik")
    form = {"stage_L1": "fikir", "stage_L2": "gelistirme", "stage_L3": "canli", "stage_L4": "atla", "stage_L5": "analiz",
            "unit_YP Lojistik": str(yp), "unit_Yedek Parça Lojistik": str(yp), "unit_VW Binek": "marka", "update_existing": "1"}
    client.post("/ayarlar/trello/aktar", data=form)
    p = db.row("SELECT * FROM projects WHERE code = 'UC88'")
    assert p["stage"] == "gelistirme" and p["department_id"] == yp and p["start_date"] == "2026-10-05"
    assert p["target_date"] == "2027-01-08" and p["health"] == "sari" and "Dependent" in p["tags"] and p["owner"] == "Burak Özkarabekir"
    assert db.scalar("SELECT code FROM projects WHERE trello_id = '68000000a2'") == "UC88-2"
    f2 = db.row("SELECT * FROM projects WHERE code = 'UC57-F2'")
    assert org.brands_by_project()[f2["id"]] == ["VW Binek"] and f2["golive_date"] == "2026-04-06"
    assert f2["department_id"] == _dept("VW Binek")
    assert not db.scalar("SELECT 1 FROM projects WHERE name = 'AI Championlar'")
    assert db.scalar("SELECT COUNT(*) FROM comments WHERE target_id = ? AND kind = 'trello'", (p["id"],)) == 1
    assert db.scalar("SELECT COUNT(*) FROM actions WHERE project_id = ? AND source = 'Trello checklist'", (p["id"],)) == 1
    # yeniden aktarım tekrar kayıt oluşturmaz
    n = db.scalar("SELECT COUNT(*) FROM projects")
    st = trello.import_board(raw, {k[6:]: v for k, v in form.items() if k.startswith("stage_")}, {}, 1)
    assert st["created"] == 0 and db.scalar("SELECT COUNT(*) FROM projects") == n
    assert db.scalar("SELECT COUNT(*) FROM comments WHERE target_id = ? AND kind = 'trello'", (p["id"],)) == 1


def test_word_table_report_links_by_uc_code(client):
    d = docx.Document()
    d.add_heading("AI IMPACT Proje Statü Raporu", 0)
    t = d.add_table(rows=2, cols=3)
    for j, v in enumerate(("Kod", "Durum", "Sonraki adım")):
        t.cell(0, j).text = v
    for j, v in enumerate(("UC 88", "Veri erişimi bekleniyor, gecikme var", "UAT")):
        t.cell(1, j).text = v
    buf = io.BytesIO()
    d.save(buf)
    r = client.post("/dokumanlar", files=[("files", ("Statü.docx", buf.getvalue(), "application/octet-stream"))],
                    data={"doc_type": "dt_durum"}, follow_redirects=False)
    did = int(r.headers["location"].rsplit("/", 1)[1])
    pipeline.drain()
    pid = db.scalar("SELECT id FROM projects WHERE code = 'UC88'")
    assert db.scalar("SELECT 1 FROM document_links WHERE document_id = ? AND project_id = ?", (did, pid))
    assert db.scalar("SELECT value_text FROM facts WHERE document_id = ? AND project_id = ? AND fact_type = 'durum'", (did, pid)).startswith("Veri erişimi")
    assert "AI bağlı değil" in client.get(f"/dokumanlar/{did}").text  # mock modda kullanıcı uyarılır


def test_edit_endpoints(client):
    pid = db.scalar("SELECT id FROM projects WHERE code = 'UC88'")
    client.post(f"/projeler/{pid}/aksiyon", data={"title": "Eski", "owner": "", "due_date": ""})
    aid = db.scalar("SELECT id FROM actions WHERE title = 'Eski'")
    client.post(f"/aksiyon/{aid}/duzenle", data={"title": "Yeni başlık", "owner": "BT", "due_date": "2026-12-01"})
    assert db.row("SELECT * FROM actions WHERE id = ?", (aid,))["title"] == "Yeni başlık"
    client.post(f"/aksiyon/{aid}/sil")
    assert not db.row("SELECT 1 FROM actions WHERE id = ?", (aid,))
    client.post(f"/projeler/{pid}/kpi", data={"name": "Stok", "unit": "", "direction": "up"})
    kid = db.scalar("SELECT id FROM kpis WHERE project_id = ? AND name = 'Stok'", (pid,))
    client.post(f"/kpi/{kid}", data={"name": "Stok devir hızı", "unit": "gün", "direction": "down", "baseline": "40"})
    k = db.row("SELECT * FROM kpis WHERE id = ?", (kid,))
    assert (k["name"], k["unit"], k["direction"], k["baseline"]) == ("Stok devir hızı", "gün", "down", 40)
    form = {"name": "Spare Parts", "code": "uc88x", "tags": "AI Impact"}
    client.post(f"/projeler/{pid}", data=form)
    assert db.scalar("SELECT code FROM projects WHERE id = ?", (pid,)) == "UC88X"
    f = next(f for f in knowledge.pending() if f["fact_type"] == "saglik") if knowledge.pending() else None
    if f:
        client.post(f"/bilgi/{f['id']}/deger", data={"value": "sari", "as_of": "2026-10-16"})
        assert db.scalar("SELECT value_text FROM facts WHERE id = ?", (f["id"],)) == "sari"
    client.post("/guvenlik/metrik", data={"period": "2026-09", "dlp_incidents": "5"})
    client.post("/guvenlik/metrik/sil", data={"period": "2026-09"})
    assert not db.scalar("SELECT 1 FROM security_metrics WHERE period = '2026-09'")
    for url in ["/ayarlar", "/guvenlik?donem=2026-10", f"/projeler/{pid}?tab=aksiyon", f"/projeler/{pid}/duzenle", "/inceleme", "/aksiyonlar"]:
        assert client.get(url).status_code == 200, url


def test_provider_autodetect(monkeypatch):
    import importlib
    monkeypatch.setenv("AI_PROVIDER", "mock")
    monkeypatch.setenv("AI_BASE_URL", "https://ai.sirket.local/v1/chat/completions")
    importlib.reload(config)
    try:
        assert config.AI_PROVIDER == "openai" and config.AI_BASE_URL == "https://ai.sirket.local/v1"
        monkeypatch.setenv("AI_BASE_URL", "https://x.openai.azure.com")
        importlib.reload(config)
        assert config.AI_PROVIDER == "azure"
    finally:
        monkeypatch.setenv("AI_PROVIDER", "mock")
        monkeypatch.setenv("AI_BASE_URL", "")
        importlib.reload(config)


def test_azure_url_variants(monkeypatch):
    """Azure portalından yapıştırılan farklı adres biçimleri doğru isteğe dönüşür."""
    import importlib
    seen = []

    class H(BaseHTTPRequestHandler):
        def do_POST(self):
            body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            seen.append((self.path, self.headers.get("api-key"), body.get("model"), "max_tokens" in body))
            ok_paths = ("/openai/deployments/gpt4o-prod/chat/completions?api-version=2025-01-01-preview",
                        "/openai/deployments/gpt4o-prod/chat/completions?api-version=2024-10-21", "/openai/v1/chat/completions")
            if self.path not in ok_paths:
                self.send_response(404)
                self.end_headers()
                self.wfile.write(b'{"error":{"code":"404","message": "Resource not found"}}')
                return
            if "max_tokens" in body:  # yeni modeller gibi davran
                self.send_response(400)
                self.end_headers()
                self.wfile.write(b'{"error":{"message":"Unsupported parameter: \'max_tokens\'. Use \'max_completion_tokens\' instead."}}')
                return
            out = json.dumps({"choices": [{"message": {"content": '{"durum": "hazir"}'}}]}).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(out)

        def log_message(self, *a):
            pass

    srv = HTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_port}"
    # yerel test sunucusu ".azure" içermediği için sağlayıcıyı açıkça azure veriyoruz; adres biçimi yine çözümlenir
    variants = [
        ("azure", f"{base}/openai/deployments/gpt4o-prod/chat/completions?api-version=2025-01-01-preview", "", "azure"),
        ("azure", base, "gpt4o-prod", "azure"),
        ("azure", f"{base}/openai/v1/", "gpt4o-prod", "azure"),
    ]
    monkeypatch.setenv("AI_API_KEY", "anahtar")
    try:
        for provider, url, model, expect in variants:
            monkeypatch.setenv("AI_PROVIDER", provider)
            monkeypatch.setenv("AI_BASE_URL", url)
            monkeypatch.setenv("AI_MODEL", model)
            importlib.reload(config)
            if url.endswith("/openai/v1/"):
                config.AI_PROVIDER = "azure_v1"  # gerçek *.azure.com adresinde otomatik algılanır
            assert config.AI_MODEL == "gpt4o-prod", url
            r = ai.diagnose()
            assert r["ok"], (url, r["steps"][-1])
        assert all(s[1] == "anahtar" for s in seen)  # api-key başlığı
        assert any(s[0].endswith("2025-01-01-preview") for s in seen)  # api-version adresten alındı
        assert seen[-1][2] == "gpt4o-prod" and not seen[-1][3]  # v1: model gövdede, max_completion_tokens'a düşüldü
        # gerçek Azure adlarında otomatik algılama
        for url, kind in [("https://dogus-ai.openai.azure.com", "azure"), ("https://dogus.cognitiveservices.azure.com/", "azure"),
                          ("https://dogus.openai.azure.com/openai/v1", "azure_v1"), ("https://x.services.ai.azure.com/models", "azure_inference")]:
            monkeypatch.setenv("AI_PROVIDER", "auto")
            monkeypatch.setenv("AI_BASE_URL", url)
            importlib.reload(config)
            assert config.AI_PROVIDER == kind, url
    finally:
        for k in ("AI_PROVIDER", "AI_BASE_URL", "AI_MODEL", "AI_API_KEY"):
            monkeypatch.setenv(k, "mock" if k == "AI_PROVIDER" else "")
        importlib.reload(config)
        srv.shutdown()


def test_deep_probe_finds_right_deployment(client):
    """Derin tanılama: yanlış model adıyla 404 alınırken deployment listesinden doğrusunu bulur."""
    from app import ai_probe

    class H(BaseHTTPRequestHandler):
        def _send(self, code, body):
            self.send_response(code)
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(json.dumps(body).encode())

        def do_GET(self):
            if self.path.startswith("/openai/deployments?"):
                return self._send(200, {"data": [{"id": "gpt4o-prod", "model": "gpt-4o", "status": "succeeded"}]})
            self._send(404, {"error": {"code": "404", "message": "Resource not found"}})

        def do_POST(self):
            self.rfile.read(int(self.headers["Content-Length"]))
            if self.headers.get("api-key") != "anahtar":
                return self._send(401, {"error": {"code": "401", "message": "Access denied"}})
            if "/deployments/gpt4o-prod/" in self.path:
                return self._send(200, {"choices": [{"message": {"content": "hazir"}}]})
            if "/deployments/" in self.path:
                return self._send(404, {"error": {"code": "DeploymentNotFound", "message": "The API deployment for this resource does not exist."}})
            self._send(404, {"error": {"code": "404", "message": "Resource not found"}})

        def log_message(self, *a):
            pass

    srv = HTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    old = (config.AI_API_KEY, config.AI_MODEL)
    config.AI_API_KEY, config.AI_MODEL = "anahtar", "gpt-4o"  # kullanıcı model adını yazmış
    try:
        p = ai_probe.Probe()
        p.azure(f"http://127.0.0.1:{srv.server_port}")
        report = "\n".join(p.lines)
        assert "deployment: gpt4o-prod" in report and "'gpt-4o' adında bir deployment yok" in report
        assert p.working and p.working[0]["AI_MODEL"] == "gpt4o-prod"
        assert "anahtar" not in report.replace("anahtar:", "")  # anahtar açık yazılmaz
        assert "ÇALIŞAN AYAR" in p._finish()
        assert client.post("/ayarlar/ai-derin-test").status_code == 200
    finally:
        config.AI_API_KEY, config.AI_MODEL = old
        srv.shutdown()
